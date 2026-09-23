"""Uredjeni Excel izvjestaji iz CSV datoteka jednog pokretanja.
Ovisnost: python -m pip install openpyxl
Poziv: python izradi_excel_izvjestaje.py --run-id test_YYYYMMDD_HHMMSS
"""
import argparse
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
REZ = ROOT / 'rezultati'
NAVY, BLUE, PALE = '14243A', '1676B8', 'EAF3FA'


def izradi(csv_putanja, xlsx_putanja, javni, run_id):
    with csv_putanja.open('r', encoding='utf-8-sig', newline='') as f:
        redovi = list(csv.DictReader(f, delimiter=';'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pregled registracija'
    headers = (['ID vozila', 'Klasa', 'Pročitana registarska oznaka - pseudonimizirano', 'OCR pouzdanost', 'Status', 'Vrijeme OCR obrade'] if javni else
               ['ID vozila', 'Klasa', 'Fotografija', 'OCR izvorni', 'OCR bez plave trake',
                'OCR odabrani', 'OCR pouzdanost', 'Status', 'Metoda', 'Vrijeme OCR obrade'])
    fields = (['ID_vozila', 'Klasa', 'HMAC_SHA256_OCR_rezultata', 'OCR_pouzdanost', 'Status', 'Vrijeme_OCR_obrade'] if javni else
              ['ID_vozila', 'Klasa', 'Datoteka_cropa', 'OCR_izvorni', 'OCR_bez_plave_trake',
               'OCR_odabrani', 'OCR_pouzdanost', 'Status', 'Metoda', 'Vrijeme_OCR_obrade'])
    last = get_column_letter(len(headers))
    ws.merge_cells(f'A1:{last}2')
    title = ws['A1']
    title.value = 'TEST 10  |  OCR REGISTARSKIH PLOČICA'
    title.font = Font(name='Aptos Display', size=19, bold=True, color='FFFFFF')
    title.fill = PatternFill('solid', fgColor=NAVY)
    title.alignment = Alignment(vertical='center', indent=1)
    for row in ws['A1:'+last+'2']:
        for cell in row:
            cell.fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 15
    ws.merge_cells(f'A3:{last}3')
    ws['A3'] = ('PSEUDONIMIZIRANI IZVJEŠTAJ' if javni else 'LOKALNI IZVJEŠTAJ – OCR REZULTATI') + f'    •    {run_id}'
    ws['A3'].font = Font(name='Aptos', size=11, bold=True, color=BLUE)
    ws.row_dimensions[3].height = 24
    broj = len(redovi)
    ocitano = sum(bool(r.get('HMAC_SHA256_OCR_rezultata' if javni else 'OCR_odabrani')) for r in redovi)
    provjera = sum(r.get('Status') == 'ZA_PROVJERU' for r in redovi)
    ws.merge_cells(f'A5:{last}5')
    ws['A5'] = f'OBRAĐENIH FOTOGRAFIJA: {broj}     |     OCR KANDIDATA: {ocitano}     |     RAZLIČITI OCR REZULTATI: {provjera}'
    ws['A5'].font = Font(name='Aptos', size=11, bold=True, color=NAVY)
    ws['A5'].fill = PatternFill('solid', fgColor=PALE)
    ws.row_dimensions[5].height = 28
    ws.merge_cells(f'A6:{last}6')
    ws['A6'] = 'Napomena: OCR kandidat nije potvrđena točna registracija. Vrijeme označava obradu, ne prolazak vozila.'
    ws['A6'].font = Font(name='Aptos', size=9, italic=True, color='526477')
    ws.row_dimensions[6].height = 22
    for c, header in enumerate(headers, 1):
        cell = ws.cell(8, c, header)
        cell.font = Font(name='Aptos', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=BLUE)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[8].height = 30
    for row_num, record in enumerate(redovi, 9):
        for col, field in enumerate(fields, 1):
            value = record.get(field, '')
            if field == 'ID_vozila':
                try: value = int(value)
                except (ValueError, TypeError): pass
            elif field == 'OCR_pouzdanost':
                try: value = float(value) if value else None
                except ValueError: value = None
            cell = ws.cell(row_num, col, value)
            cell.font = Font(name='Aptos', size=10, color=NAVY)
            cell.alignment = Alignment(vertical='center')
            if field == 'OCR_pouzdanost':
                cell.number_format = '0.00%'
            if field == 'Status':
                if value == 'ZA_PROVJERU':
                    cell.fill = PatternFill('solid', fgColor='FFF1CE')
                    cell.font = Font(name='Aptos', bold=True, color='8D5500')
                elif value in ('NIJE_IZDVOJENO', 'GRESKA_UCITAVANJA'):
                    cell.fill = PatternFill('solid', fgColor='FDE8E8')
                    cell.font = Font(name='Aptos', bold=True, color='A32323')
                else:
                    cell.fill = PatternFill('solid', fgColor='E6F4EA')
            if field in ('OCR_odabrani', 'HMAC_SHA256_OCR_rezultata'):
                cell.font = Font(name='Consolas', size=10, bold=not javni, color=NAVY)
        ws.row_dimensions[row_num].height = 23
    if redovi:
        table = Table(displayName='OCRPublic' if javni else 'OCRLocal', ref=f'A8:{last}{8+len(redovi)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True,
                                              showColumnStripes=False, showFirstColumn=False, showLastColumn=False)
        ws.add_table(table)
    ws.freeze_panes = 'A9'
    widths = ([14, 15, 69, 19, 24, 28] if javni else
              [14, 15, 28, 24, 27, 24, 19, 25, 23, 28])
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = '1:8'
    xlsx_putanja.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_putanja)
    print('[OK] Excel izvještaj:', xlsx_putanja)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--run-id', required=True)
    args = parser.parse_args()
    for javni, folder, prefix in [(False, 'privatno', 'registracije_lokalno'),
                                   (True, 'za_pregled_prije_objave', 'registracije_pseudonimizirano')]:
        csv_file = REZ / folder / f'{prefix}_{args.run_id}.csv'
        if not csv_file.is_file():
            print('[GRESKA] Nije pronađen CSV:', csv_file)
            raise SystemExit(1)
        izradi(csv_file, csv_file.with_suffix('.xlsx'), javni, args.run_id)
    print('[VAŽNO] Pseudonimizirani Excel nije automatski anoniman niti odobren za javnu objavu.')

if __name__ == '__main__':
    main()
